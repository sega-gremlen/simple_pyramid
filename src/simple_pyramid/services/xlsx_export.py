"""Выгрузка подневки в xlsx (формат файла — как в прежних версиях программы)."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from simple_pyramid.api.models.meterdata import MeterPointData

COLUMN_WIDTH = 20
UNIT_CAPTION = "кВт*ч"


def create_daily_workbook(points: list[MeterPointData], output_path: Path) -> None:
    """Создаёт xlsx: строки — параметры, столбцы — даты, значения в кВт·ч."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 — у нового Workbook активный лист есть всегда
    ws.title = "Sheet"

    # 1. Собираем все даты
    all_dates: set[datetime] = {
        val.value_dt
        for point in points
        for param in point.parameters_data
        for val in param.values
        if val.value_dt is not None
    }
    sorted_dates = sorted(all_dates)

    # 2. Заголовки
    headers = ["Прибор учета", "Параметр", *(d.strftime("%d.%m.%Y 00:00") for d in sorted_dates)]
    ws.append(headers)

    # 3. Форматирование первой строки (заголовки): жирный, по центру
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2  # следующая строка для "Точка учета"

    for point in points:
        point_caption = point.point_with_meter.caption
        meter_caption = point.point_with_meter.meter.caption

        # Строка с точкой учёта (вторая строка)
        cell = ws.cell(row=row_idx, column=1, value=f"Точка учета: {point_caption}")
        cell.font = Font(bold=True)  # жирный шрифт, центровка оставлена по умолчанию (лево)
        row_idx += 1

        # Параметры
        for param in point.parameters_data:
            # Формируем строку для столбца B: "caption, unit"
            param_display = f"{param.parameter.caption}, {UNIT_CAPTION}"

            value_by_date: dict[datetime, float] = {
                val.value_dt: val.value / 1000.0
                for val in param.values
                if val.value_dt is not None and val.value is not None
            }
            row_data: list[str | float] = [meter_caption, param_display]
            row_data.extend(value_by_date.get(d, "") for d in sorted_dates)
            ws.append(row_data)
            row_idx += 1

    # 4. Ширина всех столбцов = 20
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTH

    # 5. Перенос текста для столбцов A и B во всех строках, начиная с 3-й
    for row in range(3, ws.max_row + 1):
        for col in (1, 2):  # A = 1, B = 2
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    # 6. Устанавливаем высоту строк 3, 4, 5 = 60
    for row in (3, 4, 5):
        if row <= ws.max_row:  # проверяем, существует ли строка
            ws.row_dimensions[row].height = 60

    wb.save(output_path)
