"""Сценарии MeterService поверх эмулятора Пирамиды."""

import asyncio
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from simple_pyramid.api.errors import AmbiguousMeterError, MeterNotFoundError, PyramidError


def parts(sink, section):
    return [p for p in sink.of("MeterPart") if p["section"] == section]


async def test_search_streams_every_section(service, sink):
    session = await service.find("021250611469")
    assert sink.of("MeterPart")[0]["section"] == "identity"  # основа приходит первой
    await service.load_details(session)

    got = {p["section"] for p in sink.of("MeterPart")}
    assert got == {"identity", "info", "readings", "load", "events"}
    assert all(p["error"] is None for p in sink.of("MeterPart"))

    readings = parts(sink, "readings")[0]["data"]
    assert readings["takenAt"]
    assert readings["aPlus"]["t0"] == pytest.approx(
        readings["aPlus"]["t1"] + readings["aPlus"]["t2"]
    )
    assert parts(sink, "load")[0]["data"]["relay"] == "Включено"
    info = parts(sink, "info")[0]["data"]
    assert info["accountNumber"]
    assert session.events
    assert sink.of("Status")[-1]["kind"] == "ok"


async def test_search_not_found_and_ambiguous(service):
    with pytest.raises(MeterNotFoundError):
        await service.find("000123")
    with pytest.raises(AmbiguousMeterError):
        await service.find("111222")


async def test_no_readings_warns(service, sink):
    session = await service.find("2221")
    await service.load_details(session)
    data = parts(sink, "readings")[0]["data"]
    assert data["takenAt"] is None
    assert any(s["kind"] == "warn" and "ни одного опроса" in s["text"] for s in sink.of("Status"))


async def test_alarm_events(service, sink):
    session = await service.find("7771234")
    await service.load_details(session)
    events = parts(sink, "events")[0]["data"]
    types = {a["eventType"] for a in events["alarms"]}
    assert "Вскрытие элемента устройства" in types
    assert "Выполнен сброс состояний пломб" in types
    page = service.events_page(session, offset=0, limit=10, important_only=True)
    assert all(e.important for e in page.items)
    assert page.total == events["importantTotal"]


async def test_relay_off(service, sink):
    session = await service.find("5551")
    await service.load_details(session)
    load = parts(sink, "load")[0]["data"]
    assert load["relay"] == "Отключено"
    assert load["relayOn"] is False


async def test_section_error_does_not_break_others(service, sink, mock_app):
    session = await service.find("4441")
    mock_app.state.mock.by_point.pop(session.identity.point_id)  # точка «исчезла»
    await service.load_details(session)
    errors = [p for p in sink.of("MeterPart") if p["error"]]
    assert {p["section"] for p in errors} == {"info", "readings", "load", "events"}
    assert all("ошибку 404" in p["error"] for p in errors)


async def test_refresh_readings_updates_values(service, sink):
    session = await service.find("021250611469")
    await service.load_details(session)
    before = session.readings
    sink.events.clear()

    await service.refresh_readings(session)

    assert sink.of("SectionBusy") == [
        {"section": "readings", "busy": True},
        {"section": "readings", "busy": False},
    ]
    assert session.readings is not None
    assert before is not None
    assert session.readings.taken_at > before.taken_at  # type: ignore[operator]
    assert session.readings.a_plus.t0 > before.a_plus.t0  # type: ignore[operator]
    assert sink.of("Status")[-1]["kind"] == "ok"


async def test_refresh_load_uses_operation_date(service, sink):
    session = await service.find("021250611469")
    await service.load_details(session)
    morning = session.load
    await service.refresh_load(session)
    assert session.load is not None
    assert morning is not None
    assert session.load.relay_at > morning.relay_at  # type: ignore[operator]


async def test_refresh_events_adds_new(service):
    session = await service.find("021250611469")
    await service.load_details(session)
    count = len(session.events or [])
    await service.refresh_events(session)
    assert len(session.events or []) == count + 1


async def test_refresh_with_errors_warns(service, sink, mock_app):
    session = await service.find("9991")
    mock_app.state.mock.by_point[session.identity.point_id].behavior.slow_failing_tasks = True
    # ускоряем «долгую» задачу сценария 999
    original = mock_app.state.mock.add_task

    def fast(task):
        task.duration = 0
        return original(task)

    mock_app.state.mock.add_task = fast
    await service.refresh_load(session)
    assert sink.of("Status")[-1] == {"text": "Опрос реле: завершено с ошибками", "kind": "warn"}


async def test_refresh_cancel_restores_state(service, sink, mock_app):
    session = await service.find("021250611469")
    original = mock_app.state.mock.add_task

    def slow(task):
        task.duration = 60
        return original(task)

    mock_app.state.mock.add_task = slow
    task = asyncio.create_task(service.refresh_readings(session))
    await asyncio.sleep(0.1)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert sink.of("SectionBusy")[-1] == {"section": "readings", "busy": False}
    assert any(s["text"].endswith("отменено") for s in sink.of("Status"))


async def test_refresh_failure_clears_busy(service, sink, mock_app):
    session = await service.find("021250611469")
    mock_app.state.mock.by_meter.clear()
    with pytest.raises(PyramidError):
        await service.refresh_load(session)
    assert sink.of("SectionBusy")[-1] == {"section": "load", "busy": False}


async def test_daily_rows_and_exports(service, tmp_path):
    session = await service.find("021250611469")
    await service.load_details(session)
    end = date.today()
    start = end - timedelta(days=9)
    rows = await service.daily_rows(session, start, end)
    assert len(rows) == 10
    filled = [r for r in rows if r.t0 is not None]
    assert filled
    assert any(r.consumption for r in rows)

    daily = await service.export_daily(session, start, end, tmp_path)
    assert daily.name.startswith("Показания на начало суток по")
    assert load_workbook(daily).active.cell(1, 1).value == "Прибор учета"

    hourly = await service.export_hourly(session, start, end, tmp_path)
    assert hourly.name.startswith("Почасовка")
    assert "№021250611469" in hourly.name
    assert load_workbook(hourly).active.cell(1, 1).value.startswith("Эмулятор")


async def test_relogin_when_token_expires(settings):
    import httpx

    from dev.mock_server.app import create_app
    from simple_pyramid.api.client import PyramidClient

    app = create_app(latency=(0, 0), task_duration=(0, 0), token_ttl=0.2)
    async with PyramidClient(settings, transport=httpx.ASGITransport(app=app)) as c:
        await c.login("u", "p")
        await asyncio.sleep(0.3)
        # токен уже протух — клиент сам перелогинится и повторит запрос
        resp = await c.get_instances("021250611469")
        assert resp.data
