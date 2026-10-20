"""Выгрузка подневки в xlsx: файл совпадает с тем, что делала старая версия."""

from openpyxl import load_workbook

from simple_pyramid.api.models.meterdata import MeterPointData
from simple_pyramid.services.xlsx_export import create_daily_workbook
from tests.conftest import DATA_DIR


def _cells(path):
    ws = load_workbook(path).active
    return [[c.value for c in row] for row in ws.iter_rows()], ws


def test_daily_workbook_matches_legacy(data, tmp_path):
    points = [MeterPointData.model_validate(p) for p in data("meter_daily_data.json")]
    out = tmp_path / "daily.xlsx"
    create_daily_workbook(points, out)

    new_cells, new_ws = _cells(out)
    old_cells, old_ws = _cells(DATA_DIR / "legacy_daily.xlsx")
    assert new_cells == old_cells
    assert new_ws.column_dimensions["A"].width == old_ws.column_dimensions["A"].width == 20
    assert new_ws.row_dimensions[3].height == 60
    assert new_ws.cell(1, 1).font.bold
    assert new_ws.cell(3, 1).alignment.wrap_text


def test_daily_workbook_empty(tmp_path):
    out = tmp_path / "empty.xlsx"
    create_daily_workbook([], out)
    cells, _ = _cells(out)
    assert cells == [["Прибор учета", "Параметр"]]
